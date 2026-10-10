"""Read-only inspection of real downloads; fetch final/empty exports from local demo.
Only generated binary export copies are written; no original tester files edited.
"""
import io
import json
import os
import sys
import urllib.request
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

sys.stdout.reconfigure(encoding='utf-8')
folder = Path(__file__).parent/'evidence'
base = 'http://127.0.0.1:5000'
opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
demo_password = os.environ.get('HOSCO_QA_DEMO_PASSWORD')
if not demo_password:
    sys.exit('Set HOSCO_QA_DEMO_PASSWORD to the existing Development DemoSeed credential; never use production credentials')
login = urllib.request.Request(base+'/api/v1/auth/login', data=json.dumps({'email':'owner@hosco.local','password':demo_password}).encode('utf-8'),headers={'Content-Type':'application/json'})
with opener.open(login,timeout=30) as response: token = json.load(response)['accessToken']
branch='fdb80dd6-cf56-9599-86f7-1e2be3386090'
jan='From=2025-12-31T17:00:00Z&To=2026-01-31T16:59:59.9999999Z'
empty='From=2029-12-31T17:00:00Z&To=2030-01-31T16:59:59.9999999Z'
for name,period,fmt in [('dashboard-january-final',jan,'xlsx'),('dashboard-january-final',jan,'pdf'),('dashboard-empty-final',empty,'xlsx'),('dashboard-empty-final',empty,'pdf')]:
    req=urllib.request.Request(base+'/api/v1/reporting/dashboard/export?'+period+'&BranchId='+branch+'&format='+fmt, headers={'Authorization':'Bearer '+token})
    with opener.open(req,timeout=30) as response: (folder/(name+'.'+fmt)).write_bytes(response.read())
result={'checks':[]}
def check(name,passed,actual): result['checks'].append({'name':name,'pass':bool(passed),'actual':actual})
for name in ('dashboard-january-final','dashboard-empty-final'):
    wb=load_workbook(folder/(name+'.xlsx'),read_only=True,data_only=False)
    ws=wb['Dashboard']
    rows=list(ws.values)
    check(name+':Unicode-branch',ws['F2'].value=='HOSCO Hồ Chí Minh',ws['F2'].value)
    check(name+':no-error-cells',all(cell.data_type!='e' for row in ws for cell in row),'no spreadsheet error cells')
    if 'january' in name:
        check('xlsx-canonical-revenue',ws['B5'].value==14022500 and ws['B5'].data_type=='n',ws['B5'].value)
        check('xlsx-canonical-AOV-margin',ws['B8'].value==150779.6 and ws['B10'].value==43.7,[ws['B8'].value,ws['B10'].value])
    else:
        check('xlsx-empty-ratios-NA',ws['B8'].value=='N/A' and ws['B10'].value=='N/A',[ws['B8'].value,ws['B10'].value])
        check('xlsx-empty-revenue-numeric-zero',ws['B5'].value==0 and ws['B5'].data_type=='n',ws['B5'].value)
    wb.close()
    pdf=PdfReader(folder/(name+'.pdf'))
    text='\n'.join(page.extract_text() for page in pdf.pages)
    check(name+':pdf-UTF8-full-branch','HOSCO Hồ Chí Minh' in text,text)
    check(name+':pdf-no-mojibake',not any(x in text for x in ('Ã','Æ','Ä')),len(pdf.pages))
    if 'empty' in name:
        check('pdf-empty-NA','AOV | N/A | VND' in text and 'Gross margin | N/A | %' in text,text)
    else:
        check('pdf-canonical-revenue-AOV','Revenue | 14022500 | VND' in text and 'AOV | 150779.6 | VND' in text,text)
        check('pdf-last-inventory-retained','HOSCO-A-SKU-002 | Demo Product 2 | 1 | 8' in text,text)
result['summary']={'passed':sum(x['pass'] for x in result['checks']), 'failed':sum(not x['pass'] for x in result['checks'])}
print(json.dumps(result,ensure_ascii=False))
sys.exit(1 if result['summary']['failed'] else 0)
